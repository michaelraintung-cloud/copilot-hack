from data import *
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.3)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def para(text='', bold=False, italic=False, size=12, align=A.CENTER, after=0):
    p = d.add_paragraph(); p.alignment = align; p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return p

para('CATATAN HARIAN', True, size=14)
p = para(bold=True, size=14); p.add_run('(').bold = True
r = p.add_run('LOGBOOK'); r.bold = r.italic = True; p.add_run(')').bold = True
for r in p.runs: r.font.size = Pt(14)
para(after=12)
para('Riset Dasar Unggulan Unsrat Klaster 2 (RDUU_K2)', True, after=24)
p = para(after=24); p.add_run().add_picture('logo_unsrat.png', width=Cm(4))
para(JUDUL, True, after=60)
for t in TIM: para(t, size=11)
para(after=60)
para('UNIVERSITAS SAM RATULANGI', True)
para('2026', True)
d.add_page_break()

widths = [Cm(1.1), Cm(3.2), Cm(3.3), Cm(8.4)]
t = d.add_table(rows=1, cols=4); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
def setcell(c, text, w, bold=False, center=False):
    c.width = w; c.text = ''
    parts = text.split('\n')
    for i, part in enumerate(parts):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if center: p.alignment = A.CENTER
        r = p.add_run(part); r.bold = bold; r.font.size = Pt(11)
for c, h, w in zip(t.rows[0].cells, ['No', 'Tanggal', 'Kegiatan', 'Uraian Kegiatan'], widths):
    setcell(c, h, w, True, True)
# repeat header
trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)
for i, (w_, k, u) in enumerate(ROWS, 1):
    row = t.add_row()
    cs = row.cells
    setcell(cs[0], str(i), widths[0], center=True); setcell(cs[1], w_, widths[1])
    setcell(cs[2], k, widths[2]); setcell(cs[3], u, widths[3])
    cp = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cp)

para(after=18)
for txt in ['Manado, Oktober 2026', 'Peneliti,']:
    q = para(txt, align=A.LEFT); q.paragraph_format.left_indent = Cm(9); q.paragraph_format.keep_with_next = True
q = para(after=36); q.paragraph_format.keep_with_next = True
q = para('Michael Ch. Raintung, SE. MM', align=A.LEFT); q.paragraph_format.left_indent = Cm(9); q.paragraph_format.keep_with_next = True
para('NIP. 197612192005011001', align=A.LEFT).paragraph_format.left_indent = Cm(9)

# page numbers
fp = s.footer.paragraphs[0]; fp.alignment = A.CENTER
r = fp.add_run()
for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
    if typ:
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), typ)
    else:
        f = OxmlElement('w:instrText'); f.text = txt
    r._r.append(f)
d.save('Logbook_RDUU_K2_2026.docx')

# PDF
from reportlab.lib.pagesizes import A4
from reportlab.platypus import KeepTogether, SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER
import xml.sax.saxutils as X
B = ParagraphStyle('b', fontName='Times-Roman', fontSize=11, leading=14)
C = ParagraphStyle('c', parent=B, alignment=TA_CENTER)
H = ParagraphStyle('h', parent=C, fontName='Times-Bold')
T14 = ParagraphStyle('t', parent=H, fontSize=14, leading=18)
def pf(canvas, doc):
    canvas.setFont('Times-Roman', 10)
    if doc.page > 1: canvas.drawCentredString(A4[0]/2, 1.2*cm, str(doc.page - 1))
doc = SimpleDocTemplate('Logbook_RDUU_K2_2026.pdf', pagesize=A4, leftMargin=2.5*cm, rightMargin=2.5*cm, topMargin=2.3*cm, bottomMargin=2.3*cm,
                        title='Logbook RDUU_K2 2026', author='Michael Ch. Raintung')
E = [Paragraph('CATATAN HARIAN', T14), Paragraph('(<i>LOGBOOK</i>)', T14), Spacer(1, 12),
     Paragraph('Riset Dasar Unggulan Unsrat Klaster 2 (RDUU_K2)', ParagraphStyle('x', parent=H, fontSize=12, leading=16)),
     Spacer(1, 24), Image('logo_unsrat.png', 4*cm, 4*cm), Spacer(1, 24),
     Paragraph(JUDUL, ParagraphStyle('j', parent=H, fontSize=11)), Spacer(1, 50)]
E += [Paragraph(x_, C) for x_ in TIM] + [Spacer(1, 60), Paragraph('UNIVERSITAS SAM RATULANGI', H), Paragraph('2026', H), PageBreak()]
data = [[Paragraph('No', H), Paragraph('Tanggal', H), Paragraph('Kegiatan', H), Paragraph('Uraian Kegiatan', H)]]
for i, (w_, k, u) in enumerate(ROWS, 1):
    data.append([Paragraph(str(i), C), Paragraph(X.escape(w_), B), Paragraph(X.escape(k), B),
                 Paragraph('<br/><br/>'.join(X.escape(z) for z in u.split('\n\n')), B)])
tb = Table(data, colWidths=[1.0*cm, 3.2*cm, 3.2*cm, 8.1*cm], repeatRows=1)
tb.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .6, colors.black), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                        ('TOPPADDING', (0,0), (-1,-1), 4), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
E += [tb, Spacer(1, 24)]
sig = ParagraphStyle('s', parent=B, leftIndent=9*cm)
E += [KeepTogether([Paragraph('Manado, Oktober 2026', sig), Paragraph('Peneliti,', sig), Spacer(1, 50),
      Paragraph('Michael Ch. Raintung, SE. MM', sig), Paragraph('NIP. 197612192005011001', sig)])]
doc.build(E, onFirstPage=pf, onLaterPages=pf)
print(len(ROWS))
