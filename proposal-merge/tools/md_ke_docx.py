#!/usr/bin/env python3
"""
md_ke_docx.py
=============
Mengubah naskah proposal berformat Markdown menjadi berkas .docx siap edit/cetak
dengan tata format akademik Indonesia:

  - Ukuran kertas A4, margin kiri 4 cm, atas 4 cm, kanan 3 cm, bawah 3 cm
  - Huruf Times New Roman 12 pt, spasi 1,5, rata kiri-kanan (justify)
  - Judul bab / sub-bab memakai gaya Heading (agar otomatis masuk daftar isi)
  - Tabel Markdown dikonversi menjadi tabel Word bergaris
  - Dukungan tebal **teks**, miring *teks*, daftar berpoin, daftar bernomor,
    kutipan blok (> ), dan penanda halaman <!-- pagebreak -->

Pemakaian:
  python3 tools/md_ke_docx.py --masuk hasil/proposal-gabungan.md --keluar hasil/proposal-gabungan.docx \
      --judul "PROPOSAL DISERTASI" --penulis "Michael Raintung"

Bila modul belum ada:  pip install --break-system-packages python-docx
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
UKURAN = Pt(12)
UKURAN_JUDUL = {1: Pt(14), 2: Pt(12), 3: Pt(12)}

POLA_TEBAL = re.compile(r"(\*\*.+?\*\*|__.+?__)")
POLA_MIRING = re.compile(r"(?<!\*)\*[^*]+?\*(?!\*)")
POLA_KODE = re.compile(r"(`[^`]+`)")
POLA_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def atur_dokumen(doc: Document) -> None:
    """Terapkan format halaman & gaya dasar akademik."""
    bagian = doc.sections[0]
    bagian.page_width = Cm(21.0)
    bagian.page_height = Cm(29.7)
    bagian.left_margin = Cm(4.0)
    bagian.top_margin = Cm(4.0)
    bagian.right_margin = Cm(3.0)
    bagian.bottom_margin = Cm(3.0)

    gaya_normal = doc.styles["Normal"]
    gaya_normal.font.name = FONT
    gaya_normal.font.size = UKURAN
    gaya_normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    gaya_normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    gaya_normal.paragraph_format.space_after = Pt(6)
    gaya_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, ukuran in UKURAN_JUDUL.items():
        gaya = doc.styles[f"Heading {level}"]
        gaya.font.name = FONT
        gaya.font.size = ukuran
        gaya.font.bold = True
        gaya.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        gaya._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        gaya.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        gaya.paragraph_format.space_before = Pt(12)
        gaya.paragraph_format.space_after = Pt(6)
        gaya.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def tambah_teks_berformat(paragraf, teks: str) -> None:
    """Tulis teks dengan penanganan **tebal**, *miring*, `kode`, dan [tautan](url)."""
    teks = POLA_LINK.sub(r"\1 (\2)", teks)
    sisa = teks
    pola_gabung = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|(?<!\*)\*[^*]+?\*(?!\*))")
    for bagian in pola_gabung.split(sisa):
        if not bagian:
            continue
        if (bagian.startswith("**") and bagian.endswith("**")) or (bagian.startswith("__") and bagian.endswith("__")):
            run = paragraf.add_run(bagian[2:-2])
            run.bold = True
        elif bagian.startswith("`") and bagian.endswith("`"):
            run = paragraf.add_run(bagian[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif bagian.startswith("*") and bagian.endswith("*") and len(bagian) > 2:
            run = paragraf.add_run(bagian[1:-1])
            run.italic = True
        else:
            paragraf.add_run(bagian)
        for run in paragraf.runs:
            run.font.name = FONT
            if run.font.size is None:
                run.font.size = UKURAN


def tambah_tabel(doc: Document, baris: list[list[str]]) -> None:
    if not baris:
        return
    kolom = max(len(b) for b in baris)
    tabel = doc.add_table(rows=0, cols=kolom)
    tabel.style = "Table Grid"
    tabel.autofit = True
    for i, isi_baris in enumerate(baris):
        sel = tabel.add_row().cells
        for j in range(kolom):
            teks = isi_baris[j] if j < len(isi_baris) else ""
            par = sel[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            tambah_teks_berformat(par, teks)
            if i == 0:
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph()


def pisah_baris_tabel(baris: str) -> list[str]:
    isi = baris.strip().strip("|")
    return [sel.strip() for sel in isi.split("|")]


def baris_pemisah_tabel(baris: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:|-]+\|?", baris.strip())) and "-" in baris


def tambah_daftar_isi(doc) -> None:
    """Sisipkan judul DAFTAR ISI + field TOC Word (diperbarui otomatis saat dibuka)."""
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("DAFTAR ISI")
    run.bold = True
    run.font.size = UKURAN_JUDUL[1]
    run.font.name = FONT

    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "TOC " + chr(92) + 'o "1-3" ' + chr(92) + "h " + chr(92) + "z " + chr(92) + "u")
    isi = OxmlElement("w:r")
    teks = OxmlElement("w:t")
    teks.text = "Daftar isi akan muncul di sini. Di Word: klik kanan → Perbarui Bidang (Update Field)."
    isi.append(teks)
    fld.append(isi)
    par._p.append(fld)

    setelan = doc.settings.element
    if setelan.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        setelan.append(el)

    doc.add_page_break()


def konversi(md: Path, docx: Path, judul: str | None, penulis: str | None, daftar_isi: bool = True) -> dict:
    doc = Document()
    atur_dokumen(doc)

    if judul:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(judul.upper())
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = FONT
    if penulis:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(penulis)
        run.font.size = UKURAN
        run.font.name = FONT

    if daftar_isi:
        tambah_daftar_isi(doc)

    baris_teks = md.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")

    # lewati judul H1 pertama bila sama dengan judul halaman judul (hindari judul ganda)
    if judul:
        for n, b in enumerate(baris_teks):
            if not b.strip():
                continue
            m0 = re.match(r"^#\s+(.*)$", b.strip())
            if m0 and m0.group(1).strip().lower() == judul.strip().lower():
                baris_teks[n] = ""
            break

    i = 0
    statistik = {"paragraf": 0, "judul": 0, "tabel": 0, "daftar": 0}
    while i < len(baris_teks):
        baris = baris_teks[i]
        telanjang = baris.strip()

        if not telanjang:
            i += 1
            continue
        if telanjang.startswith("<!-- pagebreak"):
            doc.add_page_break()
            i += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", telanjang):
            i += 1
            continue

        # tabel markdown
        if telanjang.startswith("|") and i + 1 < len(baris_teks) and baris_pemisah_tabel(baris_teks[i + 1]):
            blok = [pisah_baris_tabel(telanjang)]
            i += 2
            while i < len(baris_teks) and baris_teks[i].strip().startswith("|"):
                blok.append(pisah_baris_tabel(baris_teks[i]))
                i += 1
            tambah_tabel(doc, blok)
            statistik["tabel"] += 1
            continue

        # judul
        m = re.match(r"^(#{1,6})\s+(.*)$", telanjang)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(level=level).add_run(m.group(2).strip())
            for run in doc.paragraphs[-1].runs:
                run.font.name = FONT
                run.font.size = UKURAN_JUDUL[level]
            statistik["judul"] += 1
            i += 1
            continue

        # kutipan
        if telanjang.startswith(">"):
            isi = telanjang.lstrip("> ").strip()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(1.0)
            par.paragraph_format.right_indent = Cm(0.5)
            tambah_teks_berformat(par, isi)
            for run in par.runs:
                run.italic = True
            statistik["paragraf"] += 1
            i += 1
            continue

        # daftar berpoin / bernomor / daftar periksa
        m = re.match(r"^([-*+]|\d+\.)\s+(.*)$", telanjang)
        if m:
            isi = m.group(2)
            if isi.startswith("[ ]"):
                isi = "☐ " + isi[3:].strip()
            elif isi.lower().startswith("[x]"):
                isi = "☒ " + isi[3:].strip()
            gaya = "List Number" if m.group(1)[0].isdigit() else "List Bullet"
            par = doc.add_paragraph(style=gaya)
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            tambah_teks_berformat(par, isi)
            statistik["daftar"] += 1
            i += 1
            continue

        # paragraf biasa: gabungkan baris lanjutan
        kumpulan = [telanjang]
        i += 1
        while i < len(baris_teks):
            lanjut = baris_teks[i].strip()
            if (not lanjut or lanjut.startswith(("#", "|", ">", "<!--"))
                    or re.match(r"^([-*+]|\d+\.)\s+", lanjut)
                    or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", lanjut)):
                break
            kumpulan.append(lanjut)
            i += 1
        par = doc.add_paragraph()
        tambah_teks_berformat(par, " ".join(kumpulan))
        statistik["paragraf"] += 1

    doc.save(docx)
    return statistik


def main() -> None:
    ap = argparse.ArgumentParser(description="Konversi naskah Markdown proposal menjadi .docx berformat akademik.")
    ap.add_argument("--masuk", required=True, help="Berkas .md sumber")
    ap.add_argument("--keluar", required=True, help="Berkas .docx hasil")
    ap.add_argument("--judul", help="Judul yang dicetak di awal dokumen")
    ap.add_argument("--penulis", help="Nama penulis di bawah judul")
    ap.add_argument("--tanpa-daftar-isi", action="store_true",
                    help="jangan sisipkan halaman daftar isi otomatis")
    args = ap.parse_args()

    masuk, keluar = Path(args.masuk), Path(args.keluar)
    if not masuk.exists():
        raise SystemExit(f"Berkas sumber tidak ditemukan: {masuk}")
    keluar.parent.mkdir(parents=True, exist_ok=True)
    stat = konversi(masuk, keluar, args.judul, args.penulis, not args.tanpa_daftar_isi)
    print(f"Konversi selesai: {keluar}")
    print(f"  {stat['judul']} judul · {stat['paragraf']} paragraf · {stat['tabel']} tabel · {stat['daftar']} item daftar")
    print(f"  Ukuran berkas: {keluar.stat().st_size:,} byte")


if __name__ == "__main__":
    main()
