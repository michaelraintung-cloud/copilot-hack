#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Perkakas bersama untuk membangun dokumen Word (A4) materi Fisika Dasar I."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MERAH = RGBColor(0x8B, 0x1A, 0x1A)
MERAH2 = RGBColor(0xA5, 0x2A, 0x2A)
HITAM = RGBColor(0x1C, 0x1C, 0x1C)
ABU = RGBColor(0x55, 0x55, 0x55)
ABU_MUDA = RGBColor(0x77, 0x77, 0x77)
BIRU = RGBColor(0x1F, 0x4E, 0x79)
HIJAU = RGBColor(0x3B, 0x53, 0x2B)
EMAS = RGBColor(0x8A, 0x6D, 0x1F)

BODY = "Cambria"
HEAD = "Calibri"

AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
              "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
              "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
              "w:pPrChange")
AFTER_SHD = AFTER_PBDR[1:]
AFTER_RPR = ("w:ins", "w:del", "w:moveFrom", "w:moveTo", "w:rPrChange")

TOKEN = re.compile(r"(\*\*.+?\*\*|\*(?!\*).+?\*)")
SUB = re.compile(r"(~[^~]+~)")
SUP = re.compile(r"(\^[^\^]+\^)")


# ---------------------------------------------------------------- XML helpers
def par_shading(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.insert_element_before(shd, *AFTER_SHD)


def par_border(par, **edges):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for name in ("top", "left", "bottom", "right"):
        if name in edges:
            sz, color, space = edges[name]
            el = OxmlElement("w:" + name)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), str(space))
            el.set(qn("w:color"), color)
            pBdr.append(el)
    pPr.insert_element_before(pBdr, *AFTER_PBDR)


def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_table_pct(table, pct=5000):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(pct))
    tw.set(qn("w:type"), "pct")
    tblPr.insert_element_before(
        tw, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
        "w:tblDescription", "w:tblPrChange")


def table_borders(table, color="C9C0AF", sz=6, inside=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right", "insideH", "insideV")
    for edge in edges:
        el = OxmlElement("w:" + edge)
        if edge.startswith("inside") and not inside:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def set_cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom),
                      ("right", right)):
        el = OxmlElement("w:" + name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(
        mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def field(par, instr, placeholder="", size=9, color=ABU_MUDA):
    r = par.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "begin")
    if instr.strip().startswith("TOC"):
        el.set(qn("w:dirty"), "true")
    r._r.append(el)

    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)

    r = par.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "separate")
    r._r.append(el)

    ph = par.add_run(placeholder)
    ph.font.size = Pt(size)
    ph.font.color.rgb = color

    r = par.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "end")
    r._r.append(el)


# ------------------------------------------------------------------ teks kaya
def _emit(par, text, bold, italic, size, color):
    for piece in SUB.split(text):
        if not piece:
            continue
        sub = False
        if piece.startswith("~") and piece.endswith("~") and len(piece) > 2:
            piece, sub = piece[1:-1], True
        for frag in SUP.split(piece):
            if not frag:
                continue
            sup = False
            if frag.startswith("^") and frag.endswith("^") and len(frag) > 2:
                frag, sup = frag[1:-1], True
            run = par.add_run(frag)
            run.bold = bold or None
            run.italic = italic or None
            if sub:
                run.font.subscript = True
            if sup:
                run.font.superscript = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color


def rich(par, text, size=None, color=None, italic=False, bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        _emit(par, part, b, i, size, color)
    return par


# ------------------------------------------------------------------ dokumen
def new_document(title, subject, author, keywords="Fisika Dasar, usaha, energi, "
                 "momentum", left=2.3, right=2.0):
    doc = Document()
    cp = doc.core_properties
    cp.title = title
    cp.subject = subject
    cp.author = author
    cp.keywords = keywords
    cp.category = "Bahan Ajar"

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(left)
    sec.right_margin = Cm(right)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True

    st = doc.styles["Normal"]
    st.font.name = BODY
    st.font.size = Pt(10.5)
    st.font.color.rgb = HITAM
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), BODY)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "id-ID")
    rpr.insert_element_before(lang, *AFTER_RPR)

    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.1
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc


# ------------------------------------------------------------------ komponen
def p(doc, text, size=None, space_after=4, indent=0.0, align=None,
      italic=False, bold=False, color=None, keep=False):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.keep_with_next = keep
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rich(par, text, size=size, italic=italic, bold=bold, color=color)
    return par


def rumus(doc, text, note=None, size=10.5):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.55)
    pf.right_indent = Cm(0.3)
    pf.space_before = Pt(4)
    pf.space_after = Pt(2 if note else 7)
    pf.line_spacing = 1.0
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_shading(par, "FBF8F1")
    par_border(par, left=(12, "B8860B", 8))
    rich(par, text, size=size)
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.left_indent = Cm(0.55)
        np_.paragraph_format.space_before = Pt(0)
        np_.paragraph_format.space_after = Pt(7)
        np_.paragraph_format.line_spacing = 1.0
        rich(np_, note, size=8.6, italic=True, color=ABU)
    return par


def langkah(doc, items, start=1, size=10.3):
    n = start
    for t in items:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = Cm(0.8)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(2.5)
        pf.line_spacing = 1.1
        r = par.add_run(f"{n})")
        r.bold = True
        r.font.color.rgb = MERAH2
        r.font.size = Pt(size)
        par.add_run("\t")
        rich(par, t, size=size)
        n += 1


def poin(doc, items, indent=0.7, size=10.2):
    for t in items:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = Cm(indent + 0.4)
        pf.first_line_indent = Cm(-0.4)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.08
        r = par.add_run("•\t")
        r.font.color.rgb = MERAH2
        r.bold = True
        r.font.size = Pt(size)
        rich(par, t, size=size)


def judul_bagian(doc, text, sub=None, page_break=True):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.page_break_before = page_break
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.keep_with_next = True
    r = par.add_run(text)
    r.font.name = HEAD
    r.font.size = Pt(14.5)
    r.bold = True
    r.font.color.rgb = MERAH
    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(0)
    line.paragraph_format.space_after = Pt(4 if sub else 10)
    par_border(line, bottom=(12, "8B1A1A", 4))
    if sub:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(9)
        sp.paragraph_format.keep_with_next = True
        rich(sp, sub, size=9, italic=True, color=ABU)
    return par


def judul_item(doc, no, judul, rumus_inti, sumber):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(15)
    pf.space_after = Pt(3)
    pf.keep_with_next = True
    par_shading(par, "F1E7DA")
    par_border(par, left=(20, "8B1A1A", 7))
    r = par.add_run(f"SOAL {no}")
    r.font.name = HEAD
    r.font.size = Pt(9)
    r.bold = True
    r.font.color.rgb = MERAH2
    r = par.add_run("   " + judul)
    r.font.name = HEAD
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = MERAH

    info = doc.add_paragraph()
    info.paragraph_format.space_before = Pt(0)
    info.paragraph_format.space_after = Pt(5)
    info.paragraph_format.keep_with_next = True
    info.paragraph_format.left_indent = Cm(0.08)
    r = info.add_run("Rumus inti: ")
    r.font.size = Pt(8.6)
    r.bold = True
    r.font.color.rgb = MERAH2
    rich(info, rumus_inti, size=8.8, italic=True)
    r = info.add_run("   ·   " + sumber)
    r.font.size = Pt(8)
    r.italic = True
    r.font.color.rgb = ABU_MUDA
    return par


def tabel_lambang(doc, rows, header=("Lambang", "Keterangan", "Satuan")):
    """rows: list of (lambang, arti, satuan)."""
    if not rows:
        return
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = doc.styles["Table Grid"]
    set_table_pct(table)
    table_borders(table, color="D6CDBC", sz=4)
    set_cell_margins(table, top=30, bottom=30, left=80, right=80)
    widths = (Cm(2.4), Cm(10.1), Cm(3.6))

    head = table.add_row()
    keep_row_together(head)
    for j, txt in enumerate(header):
        cell = head.cells[j]
        cell.width = widths[j]
        cell_shading(cell, "7A1F1F")
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 0.95
        r = par.add_run(txt)
        r.font.size = Pt(8.2)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (lamb, arti, satuan) in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for j, txt in enumerate((lamb, arti, satuan)):
            cell = row.cells[j]
            cell.width = widths[j]
            if i % 2 == 1:
                cell_shading(cell, "FBF8F2")
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.0
            rich(par, txt, size=8.3)
            if j == 0:
                for r in par.runs:
                    r.bold = True
            if j == 2:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(3)
    gap.paragraph_format.line_spacing = Pt(3)


def label(doc, text, color=None, size=9, space_before=6, space_after=3):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.keep_with_next = True
    r = par.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.name = HEAD
    r.font.color.rgb = color or MERAH2
    return par


def soal_teks(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.3)
    pf.right_indent = Cm(0.15)
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.12
    par_shading(par, "F6F8FB")
    par_border(par, left=(12, "1F4E79", 8))
    r = par.add_run("Soal.  ")
    r.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = BIRU
    rich(par, text, size=10.2)
    return par


def syarat(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.12)
    pf.space_before = Pt(2)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.08
    r = par.add_run("Syarat pemakaian: ")
    r.bold = True
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = ABU
    rich(par, text, size=9, italic=True, color=ABU)
    return par


def jawaban(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.2)
    pf.right_indent = Cm(0.15)
    pf.space_before = Pt(7)
    pf.space_after = Pt(7)
    pf.line_spacing = 1.08
    par_shading(par, "EEF3E8")
    par_border(par, left=(18, "4F6B3A", 8), top=(4, "C7D3BC", 4),
               bottom=(4, "C7D3BC", 4), right=(4, "C7D3BC", 4))
    r = par.add_run("Jawaban:  ")
    r.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = HIJAU
    rich(par, text, size=10.2, bold=True)
    return par


def catatan(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.2)
    pf.space_before = Pt(3)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.08
    par_shading(par, "FAFAF4")
    par_border(par, left=(10, "B8860B", 8))
    r = par.add_run("Catatan.  ")
    r.bold = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = EMAS
    rich(par, text, size=9.1, italic=True)
    return par


def gambar(doc, path, lebar_cm, caption):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(5)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(path), width=Cm(lebar_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    rich(cap, caption, size=8.4, italic=True, color=ABU_MUDA)


def header_footer(doc, kiri, kanan, footer_kiri, first_page=True):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = first_page
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.7), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(kiri)
    r.font.name = HEAD
    r.font.size = Pt(8.3)
    r.font.color.rgb = ABU_MUDA
    r = hp.add_run("\t" + kanan)
    r.font.name = HEAD
    r.font.size = Pt(8.3)
    r.font.color.rgb = ABU_MUDA
    par_border(hp, bottom=(4, "C9C0AF", 2))

    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(2)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.7), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run(footer_kiri)
    r.font.name = HEAD
    r.font.size = Pt(8.3)
    r.font.color.rgb = ABU_MUDA
    r = fp.add_run("\tHalaman ")
    r.font.name = HEAD
    r.font.size = Pt(8.3)
    r.font.color.rgb = ABU_MUDA
    field(fp, "PAGE", "1", 8.3)
    r = fp.add_run(" dari ")
    r.font.name = HEAD
    r.font.size = Pt(8.3)
    r.font.color.rgb = ABU_MUDA
    field(fp, "NUMPAGES", "1", 8.3)
    par_border(fp, top=(4, "C9C0AF", 2))


def banner_judul(doc, baris_kecil, judul_baris, sub_baris, keterangan):
    for _ in range(2):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)

    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_pct(banner)
    cell = banner.rows[0].cells[0]
    cell_shading(cell, "8B1A1A")
    cell.text = ""

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(18)
    p1.paragraph_format.space_after = Pt(2)
    r = p1.add_run(baris_kecil)
    r.font.name = HEAD
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor(0xF6, 0xE7, 0xD0)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(6)
    r = p2.add_run(judul_baris)
    r.font.name = HEAD
    r.font.size = Pt(25)
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(18)
    r = p3.add_run(sub_baris)
    r.font.name = BODY
    r.font.size = Pt(11.5)
    r.italic = True
    r.font.color.rgb = RGBColor(0xF3, 0xDF, 0xC4)

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(8)
    gap.paragraph_format.space_after = Pt(14)
    par_border(gap, bottom=(10, "B8860B", 1))

    tbl = doc.add_table(rows=0, cols=2)
    tbl.autofit = False
    for k, v in keterangan:
        row = tbl.add_row()
        c1, c2 = row.cells
        c1.width = Cm(4.0)
        c2.width = Cm(11.8)
        c1.text = ""
        c2.text = ""
        pp = c1.paragraphs[0]
        pp.paragraph_format.space_after = Pt(6)
        r = pp.add_run(k)
        r.font.name = HEAD
        r.font.size = Pt(9.8)
        r.bold = True
        r.font.color.rgb = MERAH
        pp = c2.paragraphs[0]
        pp.paragraph_format.space_after = Pt(6)
        r = pp.add_run(": " + v)
        r.font.name = BODY
        r.font.size = Pt(9.8)


def blok_penutup_cover(doc, tanggal="Oktober 2026"):
    for _ in range(2):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
    p_ = doc.add_paragraph()
    p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_.paragraph_format.space_after = Pt(2)
    r = p_.add_run("Dosen Pengasuh Mata Kuliah")
    r.font.name = HEAD
    r.font.size = Pt(9.8)
    r.font.color.rgb = ABU
    p_ = doc.add_paragraph()
    p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_.paragraph_format.space_after = Pt(2)
    r = p_.add_run("………………………………………")
    r.font.size = Pt(10.5)
    p_ = doc.add_paragraph()
    p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_.paragraph_format.space_before = Pt(8)
    r = p_.add_run(tanggal)
    r.font.name = HEAD
    r.font.size = Pt(10.5)
    r.bold = True
    r.font.color.rgb = MERAH


def normalize_order(doc):
    """Rapikan urutan elemen OOXML agar Word tidak meminta perbaikan berkas."""
    ORDER = {
        "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
                "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
                "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
                "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
                "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
                "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                "textDirection", "textAlignment", "textboxTightWrap",
                "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
        "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
                "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
                "noProof", "snapToGrid", "vanish", "webHidden", "color",
                "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
                "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl",
                "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
        "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
                  "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
                  "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
                  "tblCellMar", "tblLook", "tblCaption", "tblDescription",
                  "tblPrChange"],
        "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore",
                 "wAfter", "cantSplit", "trHeight", "tblHeader",
                 "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"],
        "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge",
                 "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
                 "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
                 "cellDel", "cellMerge", "tcPrChange"],
    }
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def fix(el):
        tag = el.tag.replace(W, "")
        order = ORDER.get(tag)
        if not order:
            return
        kids = list(el)
        seen = {}
        for child in kids:
            local = child.tag.replace(W, "")
            seen.setdefault(local, []).append(child)
        keep = {k: v[-1] for k, v in seen.items()}
        for child in kids:
            el.remove(child)
        for _, child in sorted(
                keep.items(),
                key=lambda kv: order.index(kv[0]) if kv[0] in order else len(order)):
            el.append(child)

    for el in list(doc.element.body.iter()):
        fix(el)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer):
            try:
                for el in list(part._element.iter()):
                    fix(el)
            except Exception:
                pass
    for st in list(doc.styles.element.iter()):
        fix(st)
