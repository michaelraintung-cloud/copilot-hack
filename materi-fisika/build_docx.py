#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Membangun berkas Microsoft Word (A4) dari ringkasan HTML.

Sumber : materi-fisika/ringkasan-fisika-dasar.html
Keluaran: materi-fisika/Ringkasan-Fisika-Dasar-I.docx

Jalankan:  python3 build_docx.py
"""

import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SRC = BASE / "ringkasan-fisika-dasar.html"
OUT = BASE / "Ringkasan-Fisika-Dasar-I.docx"

MERAH = RGBColor(0x8B, 0x1A, 0x1A)
MERAH2 = RGBColor(0xA5, 0x2A, 0x2A)
ABU = RGBColor(0x44, 0x44, 0x44)
ABU_MUDA = RGBColor(0x70, 0x70, 0x70)
HITAM = RGBColor(0x1C, 0x1C, 0x1C)

BODY_FONT = "Cambria"
HEAD_FONT = "Calibri"

# ----------------------------------------------------------------------------
# utilitas XML (urutan elemen pPr harus sesuai skema OOXML)
# ----------------------------------------------------------------------------
AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
AFTER_SHD = AFTER_PBDR[1:]
AFTER_RPR = (
    "w:ins", "w:del", "w:moveFrom", "w:moveTo", "w:rPrChange",
)


def _set(el, tag, **attrs):
    child = OxmlElement(tag)
    for k, v in attrs.items():
        child.set(qn("w:" + k), str(v))
    el.append(child)
    return child


def par_shading(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.insert_element_before(shd, *AFTER_SHD)


def par_border(par, **edges):
    """edges: top/left/bottom/right = (sz_eighths, color, space_pt)"""
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for name in ("top", "left", "bottom", "right"):
        if name in edges:
            sz, color, space = edges[name]
            _set(pBdr, "w:" + name, val="single", sz=sz, space=space, color=color)
    pPr.insert_element_before(pBdr, *AFTER_PBDR)


def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        _set(mar, "w:" + name, w=val, type="dxa")
    tblPr.append(mar)


def table_borders(table, color="C9C0AF", sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _set(borders, "w:" + edge, val="single", sz=sz, space=0, color=color)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def table_full_width(table, pct=5000):
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(pct))
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblPr.append(layout)


def field_run(par, instr, placeholder="", size=9, color=ABU_MUDA):
    run = par.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    if instr.strip().startswith("TOC"):
        fld.set(qn("w:dirty"), "true")
    run._r.append(fld)

    run2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run2._r.append(it)

    run3 = par.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(sep)

    ph = par.add_run(placeholder)
    ph.font.size = Pt(size)
    ph.font.color.rgb = color

    run4 = par.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run4._r.append(end)


def set_lang(rPr, lang="id-ID"):
    lang_el = OxmlElement("w:lang")
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)
    lang_el.set(qn("w:bidi"), "ar-SA")
    rPr.insert_element_before(lang_el, *AFTER_RPR)



# ----------------------------------------------------------------------------
# normalisasi urutan elemen OOXML (mencegah Word meminta "repair")
# ----------------------------------------------------------------------------
_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
            "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
            "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
            "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
            "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
            "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
            "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
            "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
            "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
            "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
              "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
              "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
              "tblCellMar", "tblLook", "tblCaption", "tblDescription",
              "tblPrChange"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
             "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc",
             "hidden", "ins", "del", "trPrChange"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
             "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
             "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
             "tcPrChange"],
}
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _normalize_element(el):
    tag = el.tag.replace(_W, "")
    order = _ORDER.get(tag)
    if not order:
        return
    children = list(el)
    seen = {}
    for child in children:
        local = child.tag.replace(_W, "")
        if local in seen:            # buang duplikat: yang terakhir dipertahankan
            seen[local].append(child)
        else:
            seen[local] = [child]
    seen = {k: v[-1] for k, v in seen.items()}
    for child in children:
        el.remove(child)
    ranked = sorted(seen.items(), key=lambda kv: (
        order.index(kv[0]) if kv[0] in order else len(order)))
    for _, child in ranked:
        el.append(child)


def normalize_order(doc):
    body = doc.element.body
    for el in list(body.iter()):
        _normalize_element(el)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer):
            try:
                for el in list(part._element.iter()):
                    _normalize_element(el)
            except Exception:
                pass
    # daftar gaya juga
    for st in list(doc.styles.element.iter()):
        _normalize_element(st)


# ----------------------------------------------------------------------------
# dokumen dasar
# ----------------------------------------------------------------------------
def build_document():
    doc = Document()

    # --- properti dokumen ---
    cp = doc.core_properties
    cp.title = "Ringkasan Materi Fisika Dasar I — Energi, Usaha, dan Momentum"
    cp.subject = ("Energi Kinetik dan Usaha; Energi Potensial dan Kekekalan Energi; "
                  "Pusat Massa, Momentum Linear, dan Tumbukan")
    cp.author = "Dosen Pengasuh Mata Kuliah Fisika Dasar I"
    cp.keywords = "Fisika Dasar, usaha, energi, momentum, tumbukan, ringkasan, soal"
    cp.category = "Bahan Ajar"
    cp.comments = "Disusun dari tiga berkas kuliah BEK1EAB4 Fisika Dasar I."

    # --- halaman A4 ---
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True

    # --- gaya Normal ---
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = HITAM
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.12
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_lang(rpr)

    return doc


def style_heading(doc, name, size, color, bold=True, before=12, after=6,
                  font=HEAD_FONT, page_break=False, border=None, indent=0.0,
                  keep_next=True, caps=False):
    st = doc.styles[name]
    st.font.name = font
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = False
    st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    set_lang(rpr)
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    pf.keep_with_next = keep_next
    pf.page_break_before = page_break
    pf.left_indent = Cm(indent)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if caps:
        st.font.all_caps = True
    if border:
        # border dipasang per-paragraf saat menulis (lebih andal)
        pass
    return st


# ----------------------------------------------------------------------------
# penulisan inline
# ----------------------------------------------------------------------------
def add_inline(par, node, bold=False, italic=False, size=None, color=None,
               font=None, sub=False, sup=False):
    """Tulis isi node (string atau Tag) ke paragraf dengan format inline."""
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        run = par.add_run(text)
        run.bold = bold or None
        run.italic = italic or None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if font:
            run.font.name = font
        if sub:
            run.font.subscript = True
        if sup:
            run.font.superscript = True
        return

    if not isinstance(node, Tag):
        return

    name = node.name.lower()
    if name == "br":
        par.add_run().add_break()
        return
    if name in ("script", "style", "nav"):
        return

    b, i, s, sp, fnt = bold, italic, size, color, font
    if name in ("strong", "b"):
        b = True
    elif name in ("em", "i"):
        i = True
    elif name == "code":
        fnt = "Consolas"
    for child in node.children:
        sub_flag = sub or name == "sub"
        sup_flag = sup or name == "sup"
        add_inline(par, child, b, i, s, sp, fnt, sub_flag, sup_flag)


def write_rich(par, node, bold=False, italic=False, size=None, color=None):
    add_inline(par, node, bold=bold, italic=italic, size=size, color=color)


# ----------------------------------------------------------------------------
# komponen dokumen
# ----------------------------------------------------------------------------
def add_rule(doc, color="D8CFC0", sz=6, space_before=2, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    par_border(p, bottom=(sz, color, 1))


def add_paragraph(doc, node, quote=False, last_in_quote=False, first_in_quote=False):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    if quote:
        pf.left_indent = Cm(0.55)
        pf.right_indent = Cm(0.25)
        pf.space_before = Pt(4 if first_in_quote else 0)
        pf.space_after = Pt(8 if last_in_quote else 2)
        pf.line_spacing = 1.08
        par_shading(par, "FBF8F1")
        par_border(par, left=(14, "B8860B", 8))
    write_rich(par, node, size=10 if quote else None)

    # baris rumus (hanya tebal, pendek) -> tengah
    plain = node.get_text(" ", strip=True)
    if quote and len(plain) <= 55 and plain and all(
        isinstance(c, Tag) and c.name in ("strong", "b") for c in node.children
        if not (isinstance(c, NavigableString) and not str(c).strip())
    ):
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par


def add_list(doc, list_tag, level=0, quote=False):
    ordered = list_tag.name.lower() == "ol"
    if ordered and list_tag.has_attr("start"):
        try:
            counter = int(list_tag["start"])
        except ValueError:
            counter = 1
    else:
        counter = 1

    for li in list_tag.find_all("li", recursive=False):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        base_indent = 0.75 + 0.65 * level
        pf.left_indent = Cm(base_indent)
        pf.first_line_indent = Cm(-0.55)
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.1
        if quote:
            par_shading(par, "FBF8F1")
            par_border(par, left=(14, "B8860B", 8))
            pf.right_indent = Cm(0.25)

        marker = f"{counter}." if ordered else ("•" if level == 0 else "–")
        run = par.add_run(marker + "\t")
        run.font.color.rgb = MERAH2 if level == 0 else HITAM
        run.bold = ordered is False and level == 0

        # isi li tanpa sublist
        inline_nodes = [c for c in li.children
                        if not (isinstance(c, Tag) and c.name in ("ul", "ol"))]
        for child in inline_nodes:
            add_inline(par, child, size=10 if quote else None)

        for sub in li.find_all(["ul", "ol"], recursive=False):
            add_list(doc, sub, level + 1, quote=quote)

        if ordered:
            counter += 1


def add_table(doc, table_tag):
    rows = []
    header = None
    thead = table_tag.find("thead")
    if thead:
        tr = thead.find("tr")
        if tr:
            header = tr.find_all(["th", "td"])
    body = table_tag.find("tbody") or table_tag
    for tr in body.find_all("tr", recursive=False):
        cells = tr.find_all(["td", "th"], recursive=False)
        if cells:
            rows.append(cells)

    if header and all(not c.get_text(strip=True) for c in header):
        header = None

    ncols = max([len(header) if header else 0] + [len(r) for r in rows] or [1])
    if ncols == 0:
        return

    table = doc.add_table(rows=0, cols=ncols)
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_full_width(table)
    table_borders(table)
    set_cell_margins(table)

    def fill_cell(cell, nodes, bold=False, color=None, size=9.0):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.0
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if isinstance(nodes, str):
            r = par.add_run(nodes)
            r.font.size = Pt(size)
            r.bold = bold
            if color:
                r.font.color.rgb = color
            return
        for n in nodes:
            if isinstance(n, Tag) and n.name == "p":
                for c in n.children:
                    add_inline(par, c, bold=bold, size=size, color=color)
            else:
                add_inline(par, n, bold=bold, size=size, color=color)

    if header:
        cells = table.add_row().cells
        repeat_header(table.rows[0])
        cant_split(table.rows[0])
        for idx in range(ncols):
            c = cells[idx]
            cell_shading(c, "7A1F1F")
            if idx < len(header):
                fill_cell(c, header[idx].contents, bold=True,
                          color=RGBColor(0xFF, 0xFF, 0xFF), size=8.8)

    for r_i, cells in enumerate(rows):
        row = table.add_row()
        cant_split(row)
        fill = "FBF8F2" if r_i % 2 == 1 else "FFFFFF"
        for idx in range(ncols):
            cell = row.cells[idx]
            cell_shading(cell, fill)
            if idx < len(cells):
                fill_cell(cell, cells[idx].contents, size=9.0)

    # jarak setelah tabel
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(4)
    gap.paragraph_format.line_spacing = Pt(3)
    return table


def add_blockquote(doc, bq):
    children = [c for c in bq.children
                if isinstance(c, Tag) or (isinstance(c, NavigableString) and c.strip())]
    real = [c for c in children if isinstance(c, Tag)]
    for idx, child in enumerate(real):
        first = idx == 0
        last = idx == len(real) - 1
        if child.name in ("ul", "ol"):
            add_list(doc, child, 0, quote=True)
        elif child.name == "p":
            add_paragraph(doc, child, quote=True, first_in_quote=first,
                          last_in_quote=last)
        elif child.name == "table":
            add_table(doc, child)
        else:
            add_paragraph(doc, child, quote=True, first_in_quote=first,
                          last_in_quote=last)


def add_problem_heading(doc, text):
    par = doc.add_heading(level=3)
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(5)
    par.paragraph_format.keep_with_next = True
    par_shading(par, "F1E7DA")
    par_border(par, left=(22, "8B1A1A", 6))
    run = par.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = HEAD_FONT
    run.font.color.rgb = MERAH
    return par


def add_section_heading(doc, text, page_break=True):
    par = doc.add_heading(level=1)
    pf = par.paragraph_format
    pf.page_break_before = page_break
    pf.space_before = Pt(0)
    pf.space_after = Pt(10)
    pf.keep_with_next = True
    par_border(par, bottom=(16, "8B1A1A", 6))
    run = par.add_run(text)
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = HEAD_FONT
    run.font.color.rgb = MERAH
    return par


def add_sub_heading(doc, text, level=2):
    par = doc.add_heading(level=level)
    pf = par.paragraph_format
    pf.space_before = Pt(14 if level == 2 else 10)
    pf.space_after = Pt(5)
    pf.keep_with_next = True
    pf.keep_together = True
    par_border(par, left=(20, "A52A2A", 7)) if level == 2 else None
    run = par.add_run(text)
    run.bold = True
    run.font.name = HEAD_FONT
    run.font.size = Pt(13 if level == 2 else 11.5)
    run.font.color.rgb = MERAH2 if level == 2 else RGBColor(0x33, 0x33, 0x33)
    return par


# ----------------------------------------------------------------------------
# halaman judul
# ----------------------------------------------------------------------------
def build_cover(doc):
    # ruang atas
    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    # banner
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_full_width(banner)
    cell = banner.rows[0].cells[0]
    cell_shading(cell, "8B1A1A")
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(22)
    p1.paragraph_format.space_after = Pt(2)
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1.add_run("RINGKASAN MATERI")
    r.font.name = HEAD_FONT
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xF6, 0xE7, 0xD0)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    r = p2.add_run("FISIKA DASAR I")
    r.font.name = HEAD_FONT
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(26)
    r = p3.add_run("Energi Kinetik & Usaha  ·  Energi Potensial & Kekekalan Energi\n"
                   "Pusat Massa, Momentum Linear & Tumbukan")
    r.font.name = BODY_FONT
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0xF3, 0xDF, 0xC4)

    # garis emas
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(18)
    par_border(p, bottom=(12, "B8860B", 1))

    # keterangan
    items = [
        ("Mata kuliah", "BEK1EAB4 — Fisika Dasar I"),
        ("Program studi", "Manajemen Rekayasa"),
        ("Fakultas", "Rekayasa Industri"),
        ("Cakupan materi", "Pertemuan minggu ke-4, ke-5, dan ke-6"),
        ("Bentuk dokumen", "Ringkasan lengkap + 58 soal latihan beserta pembahasan"),
    ]
    tbl = doc.add_table(rows=0, cols=2)
    tbl.autofit = False
    table_full_width(tbl)
    for k, v in items:
        row = tbl.add_row()
        c1, c2 = row.cells
        c1.width = Cm(4.4)
        c2.width = Cm(11.4)
        for cell in (c1, c2):
            cell.text = ""
        p = c1.paragraphs[0]
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(k)
        r.font.name = HEAD_FONT
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = MERAH
        p = c2.paragraphs[0]
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(": " + v)
        r.font.name = BODY_FONT
        r.font.size = Pt(10.5)

    # sumber
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Disusun dari tiga berkas kuliah: ")
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = ABU
    r = p.add_run("Kinetic Energy and Work; Potential Energy and Conservation of "
                  "Energy; Center of Mass – Linear Momentum and Collision")
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = ABU
    r = p.add_run(" (Karlan Rusmana, S.Si., M.T., Telkom University).")
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = ABU

    # blok bawah
    for _ in range(3):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Dosen Pengasuh Mata Kuliah")
    r.font.name = HEAD_FONT
    r.font.size = Pt(10)
    r.font.color.rgb = ABU

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("………………………………………")
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("Oktober 2026")
    r.font.name = HEAD_FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = MERAH


def build_toc(doc):
    # sengaja bukan Heading 1 agar tidak muncul di dalam daftar isi
    par = doc.add_paragraph()
    par.paragraph_format.page_break_before = True
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(12)
    par_border(par, bottom=(16, "8B1A1A", 6))
    r = par.add_run("DAFTAR ISI")
    r.font.name = HEAD_FONT
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = MERAH

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    field_run(p, r'TOC \o "1-2" \h \z \u',
              placeholder="Daftar isi sedang disiapkan…", size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Catatan: daftar isi di atas adalah field otomatis Microsoft Word. "
                  "Bila nomor halaman belum muncul, klik di dalam daftar isi, tekan "
                  "tombol kanan mouse, lalu pilih “Update Field” → “Update entire table”.")
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = ABU_MUDA


def build_header_footer(doc):
    sec = doc.sections[0]

    # header (bukan halaman pertama)
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("Ringkasan Materi Fisika Dasar I")
    r.font.name = HEAD_FONT
    r.font.size = Pt(8.5)
    r.font.color.rgb = ABU_MUDA
    r = hp.add_run("\tBEK1EAB4 — Energi, Usaha, dan Momentum")
    r.font.name = HEAD_FONT
    r.font.size = Pt(8.5)
    r.font.color.rgb = ABU_MUDA
    par_border(hp, bottom=(4, "C9C0AF", 2))

    # footer
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(2)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("Fisika Dasar I — Ringkasan & Latihan Soal")
    r.font.name = HEAD_FONT
    r.font.size = Pt(8.5)
    r.font.color.rgb = ABU_MUDA
    r = fp.add_run("\tHalaman ")
    r.font.name = HEAD_FONT
    r.font.size = Pt(8.5)
    r.font.color.rgb = ABU_MUDA
    field_run(fp, "PAGE", "1", size=8.5)
    r = fp.add_run(" dari ")
    r.font.name = HEAD_FONT
    r.font.size = Pt(8.5)
    r.font.color.rgb = ABU_MUDA
    field_run(fp, "NUMPAGES", "1", size=8.5)
    par_border(fp, top=(4, "C9C0AF", 2))


# ----------------------------------------------------------------------------
# isi utama
# ----------------------------------------------------------------------------
def build_body(doc, soup):
    main = soup.find("main")

    # mulai dari "Kata Pengantar untuk Mahasiswa"
    start = None
    for tag in main.find_all(["h2", "h1"]):
        if "Kata Pengantar" in tag.get_text():
            start = tag
            break

    add_section_heading(doc, "PENGANTAR", page_break=True)

    nodes = []
    if start:
        nodes = [start] + list(start.find_all_next())
        nodes = [n for n in nodes if n.name in (
            "h1", "h2", "h3", "h4", "p", "ul", "ol", "table", "blockquote", "hr")]
        # buang node yang berada di dalam node lain pada daftar (hindari ganda)
        nodes = [n for n in nodes if n.find_parent(["blockquote", "li", "table"]) is None]

    first_h1 = True
    for node in nodes:
        name = node.name
        if name in ("h1", "h2", "h3", "h4"):
            text = node.get_text(" ", strip=True)
            if not text:
                continue
            if name == "h1":
                add_section_heading(doc, text, page_break=not first_h1)
                first_h1 = False
            elif name == "h2":
                add_sub_heading(doc, text, level=2)
            elif name == "h3":
                if re.match(r"^Soal\s", text):
                    add_problem_heading(doc, text)
                else:
                    add_sub_heading(doc, text, level=3)
            else:
                add_sub_heading(doc, text, level=3)
        elif name == "p":
            add_paragraph(doc, node)
        elif name in ("ul", "ol"):
            add_list(doc, node)
        elif name == "table":
            add_table(doc, node)
        elif name == "blockquote":
            add_blockquote(doc, node)
        elif name == "hr":
            continue


def main():
    html = SRC.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")

    doc = build_document()
    style_heading(doc, "Heading 1", 18, MERAH, before=0, after=10)
    style_heading(doc, "Heading 2", 13, MERAH2, before=14, after=5)
    style_heading(doc, "Heading 3", 11.5, RGBColor(0x33, 0x33, 0x33),
                  before=10, after=5)

    build_header_footer(doc)
    build_cover(doc)
    build_toc(doc)
    build_body(doc, soup)

    normalize_order(doc)

    doc.save(OUT)
    print("Tersimpan:", OUT)
    print("Ukuran :", round(OUT.stat().st_size / 1024, 1), "KB")


if __name__ == "__main__":
    main()
