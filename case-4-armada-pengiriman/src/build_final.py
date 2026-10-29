# -*- coding: utf-8 -*-
"""Membangun laporan final Studi Kasus 4 (tanpa lampiran).

Keluaran: ../output/Laporan-Final-Case-4.docx
Aturan: A4, TNR 12, spasi 1,5, margin 2,54 cm, justify, isi BAB I-VI <= 8 halaman,
nomor halaman, sampul dengan logo, tanpa lampiran.

Jalankan:  python3 build_final.py
"""
import os
import re

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx_layout import (FONT, shade, table_borders, cell_margins,
                         fixed_width, no_split, add_field)
import naskah_final as N
from measure import measure

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "output", "Laporan-Final-Case-4.docx")
LOGO = os.path.join(HERE, "logo_telkom.png")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.54)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.header_distance = Cm(1.25)
sec.footer_distance = Cm(1.25)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)


def runs_of(par, text, size=12.0, bold=False, italic=False):
    for tok in re.split(r"(\*\*.+?\*\*)", text):
        if not tok:
            continue
        b = bold
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        r = par.add_run(tok)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = b
        r.italic = italic


def add_para(text, indent=True, size=12, align="just", space_before=0, space_after=0,
             keep_next=False, line=1.5):
    par = doc.add_paragraph()
    par.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY,
                     "left": WD_ALIGN_PARAGRAPH.LEFT,
                     "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    pf = par.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Cm(1.27)
    pf.keep_with_next = keep_next
    runs_of(par, text, size=size)
    return par


def add_h1(text):
    par = add_para(text, indent=False, align="center", space_before=6, space_after=6,
                   keep_next=True)
    for r in par.runs:
        r.bold = True


def add_h2(text):
    par = add_para(text, indent=False, align="left", space_before=4, space_after=0,
                   keep_next=True)
    for r in par.runs:
        r.bold = True


def add_list(items, bullet=False):
    for i, item in enumerate(items, 1):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = par.paragraph_format
        pf.line_spacing = 1.5
        pf.left_indent = Cm(1.27)
        pf.first_line_indent = Cm(-0.63)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.tab_stops.add_tab_stop(Cm(1.27))
        runs_of(par, f"{'•' if bullet else str(i) + '.'}\t{item}")


def add_table(spec):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(spec["caption"])
    r.font.name = FONT
    r.font.size = Pt(10.5)
    r.bold = True

    headers, rows, widths = spec["headers"], spec["rows"], spec["widths"]
    size = spec.get("size", 8.0)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    cell_margins(t, top=20, bottom=20, left=80, right=80)

    hdr = t.rows[0]
    hdr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    no_split(hdr)
    for i, htext in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, "EDEDED")
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_after = Pt(0)
        runs_of(par, htext, size=size, bold=True)

    for row in rows:
        cells = t.add_row().cells
        no_split(t.rows[-1])
        strong = str(row[0]).lower().startswith(("skor akhir", "jumlah", "penilaian terbobot"))
        is_label = (str(row[0]).lower().startswith("penilaian terbobot")
                    and all(not str(v).strip() for v in row[1:]))
        if is_label:
            merged = cells[0].merge(cells[-1])
            par = merged.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            runs_of(par, str(row[0]), size=size, bold=True, italic=True)
            shade(merged, "F2F2F2")
            continue
        for i, val in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            txt = str(val)
            par.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                             if i > 0 and len(txt) <= 18 else WD_ALIGN_PARAGRAPH.LEFT)
            runs_of(par, txt, size=size, bold=strong)
    fixed_width(t, widths)

    if spec.get("note"):
        par = doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(4)
        runs_of(par, spec["note"], size=8, italic=True)


# ------------------------------------------------------------------ SAMPUL
def centre(text, size=12, bold=False, before=0, after=0, line=1.15):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.line_spacing = line
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    return par


for _ in range(2):
    centre("", after=0)
centre("ANALISIS PERMASALAHAN OPERASIONAL DISTRIBUSI", size=14, bold=True)
centre("DAN EVALUASI ALTERNATIF SOLUSI BERBASIS TIK", size=14, bold=True, after=6)
centre("Studi Kasus 4: Armada Banyak, Pengiriman Tetap Terlambat", size=12, bold=True, after=26)

par = doc.add_paragraph()
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(LOGO, width=Cm(4.4))
par.paragraph_format.space_after = Pt(26)

centre("Disusun oleh:", after=2)
centre("Joanne Abigayle Kayla Raintung", after=0)
centre("102042600132  ·  MR-50-04", after=10)
centre("Dosen Pengampu: Bela Pitria Hakim", after=30)

centre("PROGRAM STUDI MANAJEMEN REKAYASA", bold=True)
centre("FAKULTAS REKAYASA INDUSTRI", bold=True)
centre("UNIVERSITAS TELKOM BANDUNG", bold=True)
centre("2026", bold=True)

doc.add_page_break()

# ------------------------------------------------------------------ RINGKASAN
par = centre("RINGKASAN", bold=True, after=12)
for text in N.RINGKASAN:
    add_para(text, space_after=6)

doc.add_page_break()

# ------------------------------------------------------------------ ISI
for blk in N.B:
    kind = blk[0]
    if kind == "p":
        add_para(blk[1], indent=blk[2], size=blk[3])
    elif kind == "h1":
        add_h1(blk[1])
    elif kind == "h2":
        add_h2(blk[1])
    elif kind == "num":
        add_list(blk[1], bullet=False)
    elif kind == "bul":
        add_list(blk[1], bullet=True)
    elif kind == "fig":
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_before = Pt(4)
        pic.paragraph_format.space_after = Pt(2)
        pic.paragraph_format.keep_with_next = True
        pic.add_run().add_picture(os.path.join(HERE, blk[1]), width=Cm(blk[2]))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(0)
        cap.paragraph_format.space_after = Pt(6)
        rr = cap.add_run(blk[4])
        rr.font.name = FONT
        rr.font.size = Pt(10.5)
        rr.bold = True
    elif kind == "tbl":
        add_table(blk[1])

# ------------------------------------------------------------------ PUSTAKA
add_h1("DAFTAR PUSTAKA")
for item in N.PUSTAKA:
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = par.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Cm(1.27)
    pf.first_line_indent = Cm(-1.27)
    pf.space_after = Pt(6)
    runs_of(par, item)

# ------------------------------------------------------------------ FOOTER
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, "PAGE", size=11)

doc.save(OUT)

body, total = measure()
print("tersimpan:", os.path.abspath(OUT))
print(f"estimasi: isi BAB I-VI = {body:.2f} hlm (batas 8) | total dokumen {total:.2f} hlm")
