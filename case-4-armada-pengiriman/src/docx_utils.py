"""Utilitas pembentuk dokumen DOCX sesuai ketentuan penulisan studi kasus:
A4, Times New Roman 12 pt, spasi 1,5, margin 2,54 cm, rata kiri-kanan.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
DARK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
FONT = "Times New Roman"


def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_borders(table, color="8EAADB", size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(el)
    return run


class Report:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()
        self._setup_page()
        self._setup_footer()
        self.log = []
        self.tbl_no = 0
        self.fig_no = {}
        self.chapter_tbl = {}
        self.chapter_fig = {}

    # ---------------------------------------------------------------- setup
    def _setup_styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for name, size, before, after, bold, color in [
            ("Heading 1", 14, 18, 10, True, NAVY),
            ("Heading 2", 12, 12, 6, True, NAVY),
            ("Heading 3", 12, 10, 4, True, DARK),
        ]:
            s = self.doc.styles[name]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.bold = bold
            s.font.color.rgb = color
            s.font.all_caps = (name == "Heading 1")
            s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _setup_page(self):
        for s in self.doc.sections:
            s.page_width = Cm(21.0)
            s.page_height = Cm(29.7)
            s.top_margin = Cm(2.54)
            s.bottom_margin = Cm(2.54)
            s.left_margin = Cm(2.54)
            s.right_margin = Cm(2.54)
            s.header_distance = Cm(1.27)
            s.footer_distance = Cm(1.27)

    def _setup_footer(self):
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = p.add_run()
        r.font.name = FONT
        r.font.size = Pt(10)
        add_field(p, "PAGE")
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(10)

    def page_break(self):
        self.log.append({"t": "pb"})
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- teks
    def h1(self, text, page_break=True):
        if page_break:
            self.page_break()
        self.log.append({"t": "h", "level": 1, "text": text})
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.log.append({"t": "h", "level": 2, "text": text})
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        self.log.append({"t": "h", "level": 3, "text": text})
        return self.doc.add_heading(text, level=3)

    def p(self, text, first_line_indent=True, spacing=1.5, after=0, align="justify",
          bold=False, italic=False, size=12, color=None, space_before=0):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(1.27) if first_line_indent else Cm(0)
        pf.space_after = Pt(after)
        pf.space_before = Pt(space_before)
        pf.line_spacing = spacing
        pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        self._rich(par, text, size=size, bold=bold, italic=italic, color=color)
        self.log.append({"t": "p", "text": text, "align": align, "indent": first_line_indent,
                         "size": size, "spacing": spacing, "after": after,
                         "space_before": space_before,
                         "color": "navy" if color is not None else None})
        return par

    def _rich(self, par, text, size=12, bold=False, italic=False, color=None):
        """Mendukung penanda **tebal** dan *miring* di dalam teks."""
        import re
        tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*)", text)
        for tok in tokens:
            if not tok:
                continue
            b, i = bold, italic
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                tok, b = tok[2:-2], True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, i = tok[1:-1], True
            r = par.add_run(tok)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = b
            r.italic = i
            if color is not None:
                r.font.color.rgb = color
        return par

    def bullets(self, items, symbol="\u2022", size=12, indent=0.0, spacing=1.5):
        for it in items:
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.75 + indent)
            pf.first_line_indent = Cm(-0.4)
            pf.line_spacing = spacing
            pf.space_after = Pt(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(f"{symbol}\t")
            r.font.name = FONT
            r.font.size = Pt(size)
            self._rich(par, it, size=size)
            self.log.append({"t": "list", "text": it, "marker": symbol, "size": size})

    def numbered(self, items, size=12, spacing=1.5, start=1):
        for i, it in enumerate(items, start):
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.9)
            pf.line_spacing = spacing
            pf.space_after = Pt(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = par.add_run(f"{i}.\t")
            r.font.name = FONT
            r.font.size = Pt(size)
            self._rich(par, it, size=size)
            self.log.append({"t": "list", "text": it, "marker": f"{i}.", "size": size})

    # ---------------------------------------------------------------- tabel
    def caption(self, text, bold=True, size=10, space_before=10, space_after=4, log=True):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.0
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.keep_with_next = True
        r = par.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = DARK
        if log:
            self.log.append({"t": "caption", "text": text, "size": size})
        return par

    def table(self, headers, rows, widths, chapter, title, note=None,
              align=None, font_size=8.5, header_size=8.5, shade_rows=True,
              first_col_bold=False, bold_rows=(), row_group_idx=None):
        self.chapter_tbl[chapter] = self.chapter_tbl.get(chapter, 0) + 1
        num = f"Tabel {chapter}.{self.chapter_tbl[chapter]}"
        self.log.append({"t": "table", "headers": headers, "rows": [[str(c) for c in r] for r in rows],
                         "widths": widths, "align": align, "font_size": font_size,
                         "header_size": header_size, "bold_rows": list(bold_rows),
                         "first_col_bold": first_col_bold, "note": note,
                         "title": f"{num}  {title}"})
        self.caption(f"{num}  {title}", log=False)

        n = len(headers)
        table = self.doc.add_table(rows=1, cols=n)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_borders(table)
        set_cell_margins(table)

        align = align or ["center"] * n
        hdr = table.rows[0]
        set_repeat_header(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(h)
            r.font.name = FONT
            r.font.size = Pt(header_size)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_shading(cell, "1F3864")

        for ri, row in enumerate(rows):
            cells = table.add_row().cells
            set_cant_split(table.rows[-1])
            for i, val in enumerate(row):
                cell = cells[i]
                cell.width = Cm(widths[i])
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                                 "left": WD_ALIGN_PARAGRAPH.LEFT,
                                 "right": WD_ALIGN_PARAGRAPH.RIGHT,
                                 "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align[i]]
                is_bold = (first_col_bold and i == 0) or (ri in bold_rows)
                self._rich(par, str(val), size=font_size, bold=is_bold)
                if shade_rows and ri % 2 == 1:
                    set_cell_shading(cell, "EDF2F9")
                if ri in bold_rows:
                    set_cell_shading(cell, "DCE6F1")

        if note:
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = Pt(3)
            pf.space_after = Pt(10)
            pf.left_indent = Cm(0.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = par.add_run(note)
            r.font.name = FONT
            r.font.size = Pt(8)
            r.italic = True
            r.font.color.rgb = GREY
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return table

    # ---------------------------------------------------------------- gambar
    def figure(self, path, width_cm, chapter, title, note=None):
        self.chapter_fig[chapter] = self.chapter_fig.get(chapter, 0) + 1
        num = f"Gambar {chapter}.{self.chapter_fig[chapter]}"
        self.log.append({"t": "figure", "path": path, "width_cm": width_cm,
                         "title": f"{num}  {title}", "note": note})
        par = self.doc.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.caption(f"{num}  {title}", space_before=0, space_after=4, log=False)
        if note:
            par = self.doc.add_paragraph()
            pf = par.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = Pt(0)
            pf.space_after = Pt(10)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = par.add_run(note)
            r.font.name = FONT
            r.font.size = Pt(8)
            r.italic = True
            r.font.color.rgb = GREY
        return num

    def save(self, path):
        self.doc.save(path)
        import json, os
        with open(os.path.splitext(path)[0] + ".layout.json", "w", encoding="utf-8") as f:
            json.dump(self.log, f, ensure_ascii=False, indent=1)
